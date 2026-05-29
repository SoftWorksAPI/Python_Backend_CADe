"""
Gerador de relatorio XLSX (Memorial Descritivo) via IA.

Gera um arquivo .xlsx com 15 abas seguindo a estrutura dos templates
de memoriais descritivos de construcao civil. Os dados sao preenchidos
por IA (OpenRouter) com base no memorial_descritivo e dados_extracao.
"""
from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from src.logger import log

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.api.v1.services.ai.client import chamar_openrouter
from src.api.v1.services.ai.prompts import (
    SYSTEM_PROMPT_RELATORIO_XLSX,
    SYSTEM_PROMPT_REVISAO_RELATORIO,
    build_xlsx_revisao_prompt,
)

OUTPUT_DIR = Path(__file__).parent / "output"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

# ---------------------------------------------------------------------------
# Estilos
# ---------------------------------------------------------------------------
NAVY_FILL = PatternFill(start_color="1A3C5E", end_color="1A3C5E", fill_type="solid")
LIGHT_NAVY_FILL = PatternFill(start_color="2C5F8A", end_color="2C5F8A", fill_type="solid")
LIGHT_GRAY_FILL = PatternFill(start_color="E8E8E8", end_color="E8E8E8", fill_type="solid")
WHITE_FONT = Font(color="FFFFFF", bold=True, size=11)
WHITE_FONT_BIG = Font(color="FFFFFF", bold=True, size=14)
BOLD_FONT = Font(bold=True, size=10)
NORMAL_FONT = Font(size=10)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

# ---------------------------------------------------------------------------
# Estrutura das 15 abas (extraida dos templates .txt)
# ---------------------------------------------------------------------------
SHEETS: list[dict[str, Any]] = [
    {
        "name": "Levantamento de campo",
        "title": "LEVANTAMENTO DE CAMPO",
        "sections": [
            {
                "name": "Dimensoes / Vaos / Alvenarias Adicionais",
                "columns": [
                    "Ambiente", "C [m]", "L [m]", "h [m]", "e [m]", "A [m2]",
                    "Tipo", "C [m]", "h [m]", "e [m]", "A [m2]", "Tipo",
                    "C [m]", "h [m]", "e [m]", "V [m3]",
                ],
                "num_rows": 20,
            },
            {
                "name": "Elementos Estruturais (Pilares, Vigas, Lajes)",
                "columns": [
                    "Ambiente", "C [m]", "L [m]", "h [m]", "A [m2]",
                    "C [m]", "h [m]", "e [m]", "A [m2]",
                    "C [m]", "L [m]", "e [m]", "A [m2]",
                ],
                "num_rows": 20,
            },
            {
                "name": "Instalacoes Eletricas",
                "columns": [
                    "Ambiente", "Quadros [Un]", "Conduletes [Un]", "Tomadas [Un]",
                    "Interruptor [Un]", "Luminarias [Un]", "Dutos [m]", "Cabos [m]",
                    "Tipo", "Acessorios [Un]", "Tipo", "Equipamentos [Un]",
                ],
                "num_rows": 20,
            },
            {
                "name": "Instalacoes Hidraulicas",
                "columns": [
                    "Ambiente", "Cavalete [Un]", "Reservatorio [L]", "Registros [Un]",
                    "Valvulas [Un]", "Torneiras [Un]", "Dutos [m]", "Tipo",
                    "Calhas [m]", "Dutos [m]", "Caixas [Un]", "Drenos [Un]",
                    "Dutos [m]", "Caixas [Un]",
                ],
                "num_rows": 20,
            },
            {
                "name": "Rede e SPDA",
                "columns": [
                    "Ambiente", "Quadros [Un]", "Conduletes [Un]", "Tomadas [Un]",
                    "Dutos [m]", "Cabos [m]", "Captacao [Un]", "Conduletes [Un]",
                    "Aterramento [m]", "Dutos [m]", "Cabos [m]",
                ],
                "num_rows": 20,
            },
            {
                "name": "Contra Incendio e Pressurizadas",
                "columns": [
                    "Ambiente", "Reservatorio [L]", "Registros [Un]", "Valvulas [Un]",
                    "Dutos [m]", "Hidrantes [Un]", "Reservatorio [L]",
                    "Registros [Un]", "Valvulas [Un]", "Dutos [m]", "Regulador [Un]",
                ],
                "num_rows": 20,
            },
            {
                "name": "Telhados",
                "columns": [
                    "Ambiente", "C [m]", "L [m]", "h [m]", "Tipo",
                    "L [m]", "C [m]", "e [m]", "A [m2]",
                ],
                "num_rows": 20,
            },
        ],
    },
    {
        "name": "Servicos Preliminares",
        "title": "1. SERVICOS PRELIMINARES",
        "sections": [
            {
                "name": "1.1 Limpeza Inicial",
                "columns": ["Descricao", "A [m2]"],
                "num_rows": 5,
            },
            {
                "name": "1.6 Interdicoes no Local da Obra e Adjacencias",
                "columns": ["Fase Ambiental", "Local", "A [m2]"],
                "num_rows": 10,
            },
            {
                "name": "1.7 Remocoes / Recolocacoes",
                "columns": [
                    "Ambiente", "Condulete", "Tomadas", "Interruptor",
                    "Luminaria", "Dutos [m]", "Cabos [m]", "Captacao",
                    "Aterr/o", "Quadros", "Postes", "Cavalete",
                    "Reservatorio", "Registros", "Valvulas", "Torneiras",
                    "Dutos [m]", "Calhas [m]", "Caixas", "Drenos",
                    "Portas", "Janelas", "Telha", "Trama", "Tesoura",
                    "Item", "Qtd",
                ],
                "num_rows": 20,
            },
            {
                "name": "1.8 Demolicoes",
                "columns": [
                    "Ambiente", "Piso [m2]", "Rodape [m]", "Azulejo [m2]",
                    "Forro [m2]", "Alvenaria Tipo", "V [m3]",
                    "Fundacao [m3]", "Pilar [m3]", "Viga [m3]", "Laje [m2]",
                ],
                "num_rows": 20,
            },
            {
                "name": "1.9 Retiras e Descarte de Residuos",
                "columns": ["Tipo", "V [m3]"],
                "num_rows": 5,
            },
        ],
    },
    {
        "name": "Movimento de Solo",
        "title": "2. MOVIMENTO DE SOLO",
        "sections": [
            {
                "name": "2.1 Terraplenagem",
                "columns": ["Descricao", "A [m2]", "V [m3]"],
                "num_rows": 10,
            },
            {
                "name": "2.2 Escavacoes",
                "columns": [
                    "Ambiente", "Tipo", "i [%]", "L [m]", "C [m]",
                    "h [m]", "Lastro", "A [m2]", "V [m3]",
                ],
                "num_rows": 20,
            },
            {
                "name": "2.3 Aterros e Reaterros",
                "columns": [
                    "Ambiente", "i [%]", "L [m]", "C [m]", "h [m]",
                    "A [m2]", "V [m3]",
                ],
                "num_rows": 20,
            },
            {
                "name": "2.7 Nivelamentos e Compactacoes",
                "columns": [
                    "Ambiente", "i [%]", "L [m]", "C [m]", "h [m]",
                    "A [m2]", "V [m3]",
                ],
                "num_rows": 20,
            },
        ],
    },
    {
        "name": "Estruturas",
        "title": "3. ESTRUTURAS",
        "sections": [
            {
                "name": "3.1.1 Sistema Viga Baldrame",
                "columns": [
                    "Peca/Secao", "L [m]", "h [m]", "C [m]", "Lastro",
                    "Concreto [m3]", "Ferragem [KgF]", "Estribo [KgF]",
                    "Forma L [m]", "Forma C [m]", "Forma A [m2]",
                ],
                "num_rows": 20,
            },
            {
                "name": "3.1.2 Estacas e Blocos de Coroamento",
                "columns": [
                    "Peca/Secao", "L [m]", "h [m]", "C [m]", "Lastro",
                    "Concreto [m3]", "Ferragem [KgF]", "Estribo [KgF]",
                    "Forma L [m]", "Forma C [m]", "Forma A [m2]",
                ],
                "num_rows": 20,
            },
            {
                "name": "3.2.1 Pilares",
                "columns": [
                    "Peca", "L [m]", "h [m]", "C [m]",
                    "Concreto [m3]", "Ferragem [KgF]", "Estribo [KgF]",
                    "Forma C [m]", "Forma L [m]", "Forma A [m2]",
                ],
                "num_rows": 20,
            },
            {
                "name": "3.2.2 Vigas",
                "columns": [
                    "Peca", "L [m]", "h [m]", "C [m]",
                    "Concreto [m3]", "Ferragem [KgF]", "Estribo [KgF]",
                    "Forma C [m]", "Forma L [m]", "Forma A [m2]",
                ],
                "num_rows": 20,
            },
            {
                "name": "3.2.3 Lajes",
                "columns": [
                    "Peca", "L [m]", "h [m]", "C [m]",
                    "Concreto [m3]", "Ferragem [KgF]", "Estribo [KgF]",
                    "Forma C [m]", "Forma L [m]", "Forma A [m2]",
                ],
                "num_rows": 20,
            },
            {
                "name": "3.3 Estruturas Metalicas",
                "columns": [
                    "Peca", "h [m]", "Perfil/Secao", "L [m]", "C [m]",
                    "Peso [KgF]", "Elast. h [m]", "Elast. C [m]",
                    "Elast. e [m]", "Elast. A [m2]", "Elast. Peso [KgF]",
                ],
                "num_rows": 20,
            },
        ],
    },
    {
        "name": "Alvenarias",
        "title": "4. ALVENARIAS",
        "sections": [
            {
                "name": "4.1 Paineis em Alvenaria",
                "columns": [
                    "Ambiente", "Peca", "C [m]", "L [m]", "h [m]",
                    "Voos", "A [m2]",
                ],
                "num_rows": 20,
            },
            {
                "name": "4.2 Vergas e Contra Vergas",
                "columns": [
                    "Ambiente", "Peca", "Qtd", "L [m]", "C [m]",
                    "h [m]", "Engastamento", "Concreto [m3]", "Ferragem [KgF]",
                ],
                "num_rows": 20,
            },
            {
                "name": "4.5 Pavimentos",
                "columns": ["Ambiente", "Local", "Tipo", "A [m2]"],
                "num_rows": 20,
            },
            {
                "name": "4.6 Cercamentos",
                "columns": [
                    "Local", "Peca", "Material", "C [m]", "h [m]",
                    "A [m2]", "Qtd",
                ],
                "num_rows": 20,
            },
            {
                "name": "4.10 Impermeabilizacao",
                "columns": [
                    "Ambiente", "Local", "Peca", "Qtd",
                    "C [m]", "L [m]", "h [m]", "ATot [m2]",
                ],
                "num_rows": 20,
            },
        ],
    },
    {
        "name": "Acabamentos",
        "title": "5. ACABAMENTOS",
        "sections": [
            {
                "name": "Pisos",
                "columns": [
                    "Ambiente", "Tipo", "e [m]", "C [m]", "L [m]", "A [m2]",
                ],
                "num_rows": 20,
            },
            {
                "name": "Rodapes",
                "columns": ["Ambiente", "Tipo", "h [m]", "C [m]", "L [m]", "A [m2]"],
                "num_rows": 20,
            },
            {
                "name": "Azulejos e Rodabancas",
                "columns": [
                    "Ambiente", "Tipo", "e [m]", "h [m]", "CTot [m]", "A [m2]",
                ],
                "num_rows": 20,
            },
            {
                "name": "Forros",
                "columns": ["Ambiente", "Tipo", "L [m]", "C [m]", "A [m2]"],
                "num_rows": 20,
            },
            {
                "name": "Pintura (Massamento, Lixamento, Selamento, Pintura)",
                "columns": [
                    "Ambiente", "h [m]", "Per [m]", "A Parede [m2]",
                    "A Teto [m2]", "Lixamento", "Selamento",
                    "Pintura Parede [m2]", "Pintura Teto [m2]",
                ],
                "num_rows": 20,
            },
            {
                "name": "Portas e Alcapoes",
                "columns": ["P/A", "Qtd", "L [cm]", "h [cm]", "e [cm]", "A [m2]"],
                "num_rows": 20,
            },
            {
                "name": "Janelas e Visores",
                "columns": ["J/V", "Qtd", "L [cm]", "h [cm]", "e [cm]", "A [m2]"],
                "num_rows": 20,
            },
            {
                "name": "Vidros",
                "columns": ["Ambiente", "L [cm]", "h [cm]", "Qtd", "A [m2]"],
                "num_rows": 20,
            },
            {
                "name": "Grades",
                "columns": ["Tipo", "A [m2]", "Malha", "e [mm]", "Afastamento [cm]"],
                "num_rows": 10,
            },
            {
                "name": "Acessorios",
                "columns": [
                    "Bacias", "Mictorios", "Lavatorios", "Cubas",
                    "Tanques", "Torneiras",
                ],
                "num_rows": 5,
            },
        ],
    },
    {
        "name": "Inst Hidraulicas",
        "title": "6. INSTALACOES HIDRAULICAS",
        "sections": [
            {
                "name": "Agua Fria",
                "columns": [
                    "Ambiente", "Cavalete [Un]", "Reservatorio [L]",
                    "Registros [Un]", "Valvulas [Un]", "Torneiras [Un]",
                    "Dutos [m]", "Tipo Duto",
                ],
                "num_rows": 20,
            },
            {
                "name": "Agua Quente",
                "columns": [
                    "Ambiente", "Aquecedor [Un]", "Registros [Un]",
                    "Valvulas [Un]", "Torneiras [Un]", "Dutos [m]", "Tipo Duto",
                ],
                "num_rows": 20,
            },
            {
                "name": "Esgoto e Pluvial",
                "columns": [
                    "Ambiente", "Dutos [m]", "Tipo Duto", "Calhas [m]",
                    "Caixas [Un]", "Drenos [Un]",
                ],
                "num_rows": 20,
            },
        ],
    },
    {
        "name": "Inst Eletricas",
        "title": "7. INSTALACOES ELETRICAS",
        "sections": [
            {
                "name": "Eletrica / SPDA / Rede",
                "columns": [
                    "Ambiente", "Quadros [Un]", "Conduletes [Un]", "Tomadas [Un]",
                    "Interruptor [Un]", "Luminarias [Un]", "Dutos [m]",
                    "Cabos [m]", "Captacao [m]", "Aterr/o [m]",
                    "Quadros Rede [Un]", "Postes [Un]",
                ],
                "num_rows": 20,
            },
        ],
    },
    {
        "name": "Inst de Telefonia e Rede",
        "title": "8. INSTALACOES DE TELEFONIA E REDE",
        "sections": [
            {
                "name": "Telefonia",
                "columns": [
                    "Ambiente", "Quadros [Un]", "Conduletes [Un]",
                    "Tomadas [Un]", "Dutos [m]", "Cabos [m]",
                ],
                "num_rows": 20,
            },
            {
                "name": "Rede de Dados",
                "columns": [
                    "Ambiente", "Quadros [Un]", "Conduletes [Un]",
                    "Tomadas [Un]", "Dutos [m]", "Cabos [m]",
                ],
                "num_rows": 20,
            },
        ],
    },
    {
        "name": "Inst Mecanicas",
        "title": "9. INSTALACOES MECANICAS",
        "sections": [
            {
                "name": "Ar Condicionado e Ventilacao",
                "columns": [
                    "Ambiente", "Local", "AC Qt [Un]", "AC Pot [BTU]",
                    "Dutos [m]", "Cabo Eletrico [m]", "Gas Refrigerante [m]",
                    "Ventilador Qt [Un]", "Ventilador Tipo",
                    "Exaustor Qt [Un]", "Exaustor Tipo",
                ],
                "num_rows": 20,
            },
        ],
    },
    {
        "name": "Inst Pressurizadas",
        "title": "10. INSTALACOES PRESSURIZADAS",
        "sections": [
            {
                "name": "Terminais e Hidrantes",
                "columns": [
                    "Ambiente", "Terminal Qt [Un]", "Terminal Dimensoes",
                    "Barra [m]", "Cordalha [m]", "Duto [m]",
                    "Hidrante Qt [Un]", "Hidrante Dimensoes",
                    "Registro [Un]", "Valvula [Un]",
                ],
                "num_rows": 20,
            },
            {
                "name": "Extintores Portateis",
                "columns": ["Local", "Tipo", "Peso [Kg]", "Capacidade", "Quantidade"],
                "num_rows": 10,
            },
        ],
    },
    {
        "name": "Inst de Seguranca",
        "title": "11. INSTALACOES DE SEGURANCA",
        "sections": [
            {
                "name": "Plaquetas de Sinalizacao",
                "columns": [
                    "Local", "Saida [Un]", "Extintor [Un]",
                    "Quadro Forca [Un]", "Hidrante [Un]", "Alarme [Un]",
                ],
                "num_rows": 20,
            },
            {
                "name": "Sinalizacao de Via",
                "columns": [
                    "Local", "Proibido Fumar", "Perigo Inflamavel",
                    "Risco Explosao", "Contra Mao", "Curva Direita",
                    "Curva Esquerda", "40 Km/h", "Pare",
                ],
                "num_rows": 20,
            },
        ],
    },
    {
        "name": "Comunicacoes Ambientais",
        "title": "12. COMUNICACOES AMBIENTAIS",
        "sections": [
            {
                "name": "Comunicacao Visual e Sonora",
                "columns": ["Ambiente", "Local", "Tipo", "Quantidade", "Observacao"],
                "num_rows": 20,
            },
        ],
    },
    {
        "name": "Paisagismos",
        "title": "13. PAISAGISMOS",
        "sections": [
            {
                "name": "Areas Verdes e Jardimizacao",
                "columns": [
                    "Local", "Tipo", "Area [m2]", "Qtd Plantas",
                    "Irrigacao", "Observacao",
                ],
                "num_rows": 20,
            },
        ],
    },
    {
        "name": "Entrega da Obra",
        "title": "14. ENTREGA DA OBRA",
        "sections": [
            {
                "name": "Documentacao e Vistoria Final",
                "columns": [
                    "Item", "Descricao", "Responsavel", "Prazo", "Status",
                ],
                "num_rows": 20,
            },
        ],
    },
]


# ---------------------------------------------------------------------------
# Prompt builder
# ---------------------------------------------------------------------------

def build_xlsx_prompt(
    memorial_descritivo: dict[str, Any],
    dados_extracao: dict[str, Any],
    normas_contexto: str = "",
) -> str:
    """Monta o user prompt para a IA preencher todas as abas do XLSX."""

    # Estrutura compacta das abas
    estrutura_texto = ""
    for sheet in SHEETS:
        estrutura_texto += f"\n## Aba: {sheet['name']}\n"
        for sec in sheet["sections"]:
            cols = ", ".join(sec["columns"])
            estrutura_texto += f"  - {sec['name']} ({sec['num_rows']} linhas): [{cols}]\n"

    # Contexto do memorial
    memorial_json = json.dumps(memorial_descritivo, ensure_ascii=False, indent=2)[:4000]
    extracao_json = json.dumps(dados_extracao, ensure_ascii=False, indent=2)[:2000]

    prompt = f"""Preencha TODAS as abas e secoes abaixo com dados realistas para o projeto descrito.

## Memorial Descritivo do Projeto
{memorial_json}

## Dados de Extracao
{extracao_json}

## Normas Tecnicas (RAG)
{normas_contexto[:1500] if normas_contexto else "Nenhuma norma disponivel."}

## Estrutura das Abas
{estrutura_texto}

## Instrucoes
- Retorne APENAS JSON valido (sem markdown fences, sem texto antes ou depois).
- Preencha TODAS as linhas de dados de TODAS as secoes com valores realistas.
- Use nomes de ambientes coerentes com o tipo de projeto do memorial.
- Use virgula como separador decimal (ex: "12,50").
- Para colunas de texto (Tipo, Peca, etc.), use termos tecnicos realistas.
- Numeros devem ser strings com virgula decimal (ex: "15,00", "3,50").
- Para colunas sem dados relevantes, use "0,00" ou "".

## Formato JSON esperado
{{
  "sheets": [
    {{
      "name": "Nome da Aba",
      "sections": [
        {{
          "section_name": "Nome da Secao",
          "rows": [
            ["valor_col1", "valor_col2", ...],
            ...
          ],
          "total_row": ["Total", "soma1", "soma2", ...]
        }}
      ]
    }}
  ]
}}"""

    return prompt


# ---------------------------------------------------------------------------
# Response parser
# ---------------------------------------------------------------------------

def parse_ai_response(texto_ia: str) -> dict:
    """Parseia a resposta JSON da IA com multiplas tentativas de fallback."""
    texto = texto_ia.strip()

    # Tentativa 1: JSON direto
    try:
        return json.loads(texto)
    except json.JSONDecodeError:
        pass

    # Tentativa 2: Extrair de code fence
    match = re.search(r"```(?:json)?\s*\n?(.*?)\n?```", texto, re.DOTALL)
    if match:
        try:
            return json.loads(match.group(1).strip())
        except json.JSONDecodeError:
            pass

    # Tentativa 3: Extrair primeiro bloco {...}
    start = texto.find("{")
    end = texto.rfind("}")
    if start != -1 and end != -1 and end > start:
        try:
            return json.loads(texto[start : end + 1])
        except json.JSONDecodeError:
            pass

    log.warn("XLSX", f"Nao foi possivel parsear resposta da IA ({len(texto)} chars)")
    return {"sheets": []}


# ---------------------------------------------------------------------------
# XLSX writer
# ---------------------------------------------------------------------------

def _auto_width(ws, min_width: int = 8, max_width: int = 30):
    """Ajusta largura das colunas baseado no conteudo."""
    for col_cells in ws.columns:
        max_len = min_width
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value:
                max_len = max(max_len, min(len(str(cell.value)) + 2, max_width))
        ws.column_dimensions[col_letter].width = max_len


def _is_numeric_str(val: str) -> bool:
    """Verifica se uma string representa um numero (formato BR: 12,50)."""
    if not isinstance(val, str):
        return False
    cleaned = val.replace(",", ".").strip()
    try:
        float(cleaned)
        return True
    except ValueError:
        return False


def _str_to_float(val: str) -> float | None:
    """Converte string BR para float."""
    if not isinstance(val, str):
        return None
    cleaned = val.replace(",", ".").strip()
    try:
        return float(cleaned)
    except ValueError:
        return None


def gerar_xlsx(ai_data: dict, arquivo_original: str = "N/A") -> str:
    """Gera o arquivo XLSX com todas as abas e retorna o path."""
    wb = Workbook()
    # Remover aba padrao
    wb.remove(wb.active)

    sheets_data = {s["name"]: s for s in ai_data.get("sheets", [])}

    for sheet_def in SHEETS:
        ws = wb.create_sheet(title=sheet_def["name"][:31])  # Excel limita a 31 chars

        # Linha 1: Titulo da aba
        row = 1
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(len(s["columns"]) for s in sheet_def["sections"]))
        title_cell = ws.cell(row=row, column=1, value=sheet_def["title"])
        title_cell.font = WHITE_FONT_BIG
        title_cell.fill = NAVY_FILL
        title_cell.alignment = HEADER_ALIGN
        row += 2  # linha em branco

        # Dados da IA para esta aba
        sheet_ai = sheets_data.get(sheet_def["name"], {})
        sections_ai = {s["section_name"]: s for s in sheet_ai.get("sections", [])}

        for sec_def in sheet_def["sections"]:
            num_cols = len(sec_def["columns"])

            # Subtitulo da secao
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=num_cols)
            sub_cell = ws.cell(row=row, column=1, value=sec_def["name"])
            sub_cell.font = WHITE_FONT
            sub_cell.fill = LIGHT_NAVY_FILL
            sub_cell.alignment = HEADER_ALIGN
            row += 1

            # Cabecalhos das colunas
            for ci, col_name in enumerate(sec_def["columns"], 1):
                cell = ws.cell(row=row, column=ci, value=col_name)
                cell.font = BOLD_FONT
                cell.fill = LIGHT_GRAY_FILL
                cell.alignment = HEADER_ALIGN
                cell.border = THIN_BORDER
            row += 1

            # Dados da IA
            sec_ai = sections_ai.get(sec_def["name"], {})
            ai_rows = sec_ai.get("rows", [])
            ai_total = sec_ai.get("total_row")

            for ri in range(sec_def["num_rows"]):
                ai_row = ai_rows[ri] if ri < len(ai_rows) else []
                for ci in range(num_cols):
                    val = ai_row[ci] if ci < len(ai_row) else ""
                    if val is None:
                        val = ""
                    cell = ws.cell(row=row, column=ci + 1, value=str(val))
                    cell.font = NORMAL_FONT
                    cell.alignment = LEFT_ALIGN
                    cell.border = THIN_BORDER
                    # Alternancia de cor
                    if ri % 2 == 1:
                        cell.fill = PatternFill(
                            start_color="F5F5F5", end_color="F5F5F5", fill_type="solid"
                        )
                row += 1

            # Linha Total
            total_label_written = False
            for ci in range(num_cols):
                col_letter = get_column_letter(ci + 1)
                # Verificar se a coluna tem dados numericos
                has_numeric = False
                if ai_rows:
                    for ai_row in ai_rows[:5]:
                        if ci < len(ai_row) and _is_numeric_str(str(ai_row[ci])):
                            has_numeric = True
                            break

                if has_numeric:
                    formula = f"=SUM({col_letter}{row - sec_def['num_rows']}:{col_letter}{row - 1})"
                    cell = ws.cell(row=row, column=ci + 1, value=formula)
                elif not total_label_written and ci == 0:
                    cell = ws.cell(row=row, column=ci + 1, value="Total")
                    total_label_written = True
                else:
                    # Se tem total_row da IA, usar
                    if ai_total and ci < len(ai_total) and ai_total[ci]:
                        cell = ws.cell(row=row, column=ci + 1, value=str(ai_total[ci]))
                    else:
                        cell = ws.cell(row=row, column=ci + 1, value="")

                cell.font = BOLD_FONT
                cell.border = Border(
                    left=Side(style="thin"),
                    right=Side(style="thin"),
                    top=Side(style="double"),
                    bottom=Side(style="thin"),
                )
            row += 2  # espaco entre secoes

        _auto_width(ws)

    # Salvar
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    nome = Path(arquivo_original).stem if arquivo_original != "N/A" else "relatorio"
    nome_arquivo = f"{nome}_memorial_{ts}.xlsx"
    caminho = OUTPUT_DIR / nome_arquivo
    wb.save(str(caminho))
    log.success("XLSX", f"Arquivo gerado: {caminho} ({caminho.stat().st_size} bytes)")
    return str(caminho)


# ---------------------------------------------------------------------------
# Orquestrador principal
# ---------------------------------------------------------------------------

def gerar_relatorio_xlsx(
    memorial_descritivo: dict[str, Any],
    dados_extracao: dict[str, Any],
    arquivo_original: str = "N/A",
    normas_contexto: str = "",
) -> tuple[str, str]:
    """
    Pipeline completo: monta prompt -> chama IA -> parseia -> gera XLSX -> revisao.

    Returns:
        Tupla (path do arquivo XLSX, revisao em Markdown).
    """
    log.info("XLSX", "Iniciando geracao de relatorio XLSX...")

    # 1. Montar prompt
    user_prompt = build_xlsx_prompt(memorial_descritivo, dados_extracao, normas_contexto)
    log.info("XLSX", f"Prompt montado ({len(user_prompt)} chars)")

    # 2. Chamar IA
    log.info("XLSX", "Chamando OpenRouter...")
    texto_ia = chamar_openrouter(SYSTEM_PROMPT_RELATORIO_XLSX, user_prompt)
    log.info("XLSX", f"IA retornou {len(texto_ia)} chars")

    # 3. Parsear resposta
    ai_data = parse_ai_response(texto_ia)
    sheets_count = len(ai_data.get("sheets", []))
    log.info("XLSX", f"Resposta parseada: {sheets_count} abas")

    # 4. Gerar XLSX
    caminho = gerar_xlsx(ai_data, arquivo_original)

    # 5. Gerar revisao da IA
    log.info("XLSX", "Gerando revisao da IA...")
    try:
        revisao_prompt = build_xlsx_revisao_prompt(memorial_descritivo, ai_data)
        revisao_md = chamar_openrouter(SYSTEM_PROMPT_REVISAO_RELATORIO, revisao_prompt)
        log.success("XLSX", f"Revisao gerada ({len(revisao_md)} chars)")
    except Exception as e:
        log.warn("XLSX", f"Falha ao gerar revisao: {e}")
        revisao_md = f"## Revisao indisponivel\n\nErro ao gerar revisao: {str(e)}"

    return caminho, revisao_md
